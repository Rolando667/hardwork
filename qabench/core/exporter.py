"""Експорт результатів у CSV та XLSX."""
from __future__ import annotations

import csv
from pathlib import Path
from typing import Any, List

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .aggregate import Aggregation
from .models import Dialog, ModelConfig


_HEADERS = [
    "Діалог",
    "Еталон",
    "Модель",
    "Рівень",
    "Вердикт",
    "Нестабільність",
    "Коментар",
    "Токени вх",
    "Токени вих",
    "Вартість $",
    "Час с",
    "Трансляція роздумів",
]


def _fraction(cell) -> str:
    if cell.n_total > 1 and cell.instable and cell.n_majority:
        return f"{cell.n_majority}/{cell.n_total}"
    return ""


def _rows(
    agg: Aggregation,
    dialogs: List[Dialog],
) -> List[List[Any]]:
    dialog_by_id = {d.id: d for d in dialogs}
    rows: List[List[Any]] = []
    for dialog in dialogs:
        for col in agg.columns:
            cell = agg.cells.get((dialog.id, col.model_id, col.level))
            if cell is None:
                continue
            verdict = cell.majority_status or (f"⚠️ {cell.error}" if cell.error else "⚠️")
            rows.append(
                [
                    dialog.name,
                    dialog.etalon,
                    col.model_name,
                    col.level,
                    verdict,
                    _fraction(cell),
                    cell.comment,
                    cell.tokens_in if cell.tokens_in is not None else "",
                    cell.tokens_out if cell.tokens_out is not None else "",
                    round(cell.cost, 6) if cell.cost is not None else "",
                    round(cell.avg_seconds, 2),
                    cell.translation,
                ]
            )
    return rows


def export_csv(path: str | Path, agg: Aggregation, dialogs: List[Dialog]) -> None:
    rows = _rows(agg, dialogs)
    with open(path, "w", newline="", encoding="utf-8-sig") as fh:
        writer = csv.writer(fh)
        writer.writerow(_HEADERS)
        writer.writerows(rows)
        # Порожній рядок + підсумки по колонках.
        writer.writerow([])
        writer.writerow(
            ["Колонка", "Збігів", "Хибних алертів", "Пропусків", "Вартість $", "Сер. час с"]
        )
        for col in agg.columns:
            s = agg.col_summary.get((col.model_id, col.level))
            if s is None:
                continue
            writer.writerow(
                [
                    col.label,
                    s.matches,
                    s.false_alerts,
                    s.misses,
                    round(s.cost, 6),
                    round(s.avg_time, 2),
                ]
            )
        writer.writerow([])
        writer.writerow(["Загальна вартість прогону, $", round(agg.total_cost, 6)])


def export_xlsx(path: str | Path, agg: Aggregation, dialogs: List[Dialog]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Результати"

    header_font = Font(bold=True)
    header_fill = PatternFill("solid", fgColor="DDDDDD")

    ws.append(_HEADERS)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill

    dialog_by_id = {d.id: d for d in dialogs}
    for dialog in dialogs:
        for col in agg.columns:
            agg_cell = agg.cells.get((dialog.id, col.model_id, col.level))
            if agg_cell is None:
                continue
            verdict = agg_cell.majority_status or (
                f"⚠️ {agg_cell.error}" if agg_cell.error else "⚠️"
            )
            ws.append(
                [
                    dialog.name,
                    dialog.etalon,
                    col.model_name,
                    col.level,
                    verdict,
                    _fraction(agg_cell),
                    agg_cell.comment,
                    agg_cell.tokens_in if agg_cell.tokens_in is not None else "",
                    agg_cell.tokens_out if agg_cell.tokens_out is not None else "",
                    round(agg_cell.cost, 6) if agg_cell.cost is not None else "",
                    round(agg_cell.avg_seconds, 2),
                    agg_cell.translation,
                ]
            )
            row_idx = ws.max_row
            if agg_cell.color == "green":
                fill = PatternFill("solid", fgColor="C6EFCE")
            elif agg_cell.color == "red":
                fill = PatternFill("solid", fgColor="FFC7CE")
            else:
                fill = None
            if fill is not None:
                ws.cell(row=row_idx, column=5).fill = fill

    _autosize(ws)

    # Аркуш підсумків.
    ws2 = wb.create_sheet("Підсумки")
    ws2.append(
        ["Колонка", "Збігів", "Хибних алертів", "Пропусків", "Вартість $", "Сер. час с"]
    )
    for cell in ws2[1]:
        cell.font = header_font
        cell.fill = header_fill
    for col in agg.columns:
        s = agg.col_summary.get((col.model_id, col.level))
        if s is None:
            continue
        ws2.append(
            [
                col.label,
                s.matches,
                s.false_alerts,
                s.misses,
                round(s.cost, 6),
                round(s.avg_time, 2),
            ]
        )
    ws2.append([])
    total_row = ["Загальна вартість прогону, $", "", "", "", round(agg.total_cost, 6), ""]
    ws2.append(total_row)
    ws2.cell(row=ws2.max_row, column=1).font = header_font
    _autosize(ws2)

    wb.save(str(path))


def _autosize(ws) -> None:
    for column_cells in ws.columns:
        length = 0
        letter = column_cells[0].column_letter
        for cell in column_cells:
            value = cell.value
            if value is None:
                continue
            length = max(length, len(str(value)))
        ws.column_dimensions[letter].width = min(60, max(10, length + 2))
